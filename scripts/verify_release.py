"""Read-only checks of release contents and sample credential/contact patterns."""

import os
import re
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from openpyxl import load_workbook

from scripts.check_delivery import TEXT_SUFFIXES, validate_delivery_text
from src.version import bundle_name

ROOT = Path(__file__).resolve().parents[1]
PATTERNS = [
    re.compile(r"sk-[A-Za-z0-9_-]{16,}"),
    re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"),
]


def main() -> None:
    with ZipFile(ROOT / "dist" / bundle_name()) as archive:
        names = archive.namelist()
        for name in names:
            if Path(name).suffix in TEXT_SUFFIXES:
                validate_delivery_text(archive.read(name).decode("utf-8"), name)
        assert not any(
            name.endswith("/.env")
            or any(
                part in name
                for part in ("/.agents/", "/node_modules/", "/__pycache__/", "/runtime/", "/.venv/")
            )
            for name in names
        )
        sample_texts = []
        for name in names:
            if "/data/" not in name:
                continue
            content = archive.read(name)
            if name.endswith((".md", ".json")):
                sample_texts.append(content.decode("utf-8"))
            elif name.endswith(".xlsx"):
                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
                try:
                    sample_texts.append(
                        "\n".join(
                            str(value)
                            for sheet in workbook
                            for row in sheet.iter_rows(values_only=True)
                            for value in row
                            if value is not None
                        )
                    )
                finally:
                    workbook.close()
        for text in sample_texts:
            assert not any(pattern.search(text) for pattern in PATTERNS), (
                "Sample matches a credential/contact pattern; inspect privately."
            )
        template = archive.read("ai-business-solution-agent/.env.example").decode()
        assert re.search(r"^OPENAI_API_KEY=\s*$", template, re.M)
        assert re.search(r"^DEEPSEEK_API_KEY=\s*$", template, re.M)
        assert re.search(r"^APP_ACCESS_PASSWORD=\s*$", template, re.M)
        assert len([name for name in names if "/screenshots/" in name and name.endswith(".png")]) == 4
        assert "ai-business-solution-agent/media/demo-local.mp4" in names
        print(
            f"PASS: {len(names)} allowlisted files, empty secret template, four screenshots/video, no sample credential/contact matches."
        )
        # Separate interpreter/cwd ensures imports come from the extracted bundle, not this worktree.
        with TemporaryDirectory(prefix="diagnosis-release-") as folder:
            archive.extractall(folder)
            root = Path(folder) / "ai-business-solution-agent"
            environment = {
                **os.environ,
                "APP_ENV": "development",
                "APP_ACCESS_PASSWORD": "",
                "OPENAI_API_KEY": "",
                "DEEPSEEK_API_KEY": "",
                "LLM_PROVIDER": "openai",
            }
            code = """
from pathlib import Path
import src
from streamlit.testing.v1 import AppTest
assert Path(src.__file__).resolve().parent == Path.cwd() / 'src'
app = AppTest.from_file('app.py').run()
app.button[0].click().run()
assert not app.exception and app.get('download_button')
app.selectbox[0].select('消费品牌团队').run()
app.button[0].click().run()
assert not app.exception and app.get('download_button')
print('PASS: extracted bundle imports its own src; both page cases and downloads work.')
"""
            subprocess.run([sys.executable, "-c", code], cwd=root, env=environment, check=True)


if __name__ == "__main__":
    main()
